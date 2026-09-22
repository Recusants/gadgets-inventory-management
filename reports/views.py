"""
Function-Based Views for Reports & Analytics with Excel and PDF Export.
Strictly zero Django Forms, 100% FBVs, offline ReportLab and openpyxl engines.
"""

from decimal import Decimal
import io
from datetime import datetime, timedelta
from django.shortcuts import render
from django.http import HttpResponse
from django.db.models import Sum, Count, Q, F
from django.utils import timezone

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from sales.models import Sale, SaleItem
from inventory.models import Product, Category, Stock, StockBatch
from expenses.models import Expense
from core.models import CompanySetting
from core.validators import _parse_custom_date


def get_date_range(filter_type, from_date_raw, to_date_raw):
    """Resolve start and end dates based on preset or custom range."""
    today = timezone.now().date()
    start_date = None
    end_date = today

    if filter_type == 'today':
        start_date = today
    elif filter_type == 'week':
        start_date = today - timedelta(days=7)
    elif filter_type == 'month':
        start_date = today - timedelta(days=30)
    elif filter_type == 'year':
        start_date = today - timedelta(days=365)
    elif filter_type == 'custom':
        if from_date_raw:
            start_date = _parse_custom_date(from_date_raw)
        if to_date_raw:
            end_date = _parse_custom_date(to_date_raw) or today

    if not start_date:
        start_date = today - timedelta(days=30)  # Default: last 30 days

    return start_date, end_date


def report_index_view(request):
    """
    Renders the Reports & Analytics hub.
    Supports interactive date range filtering, expenses factoring, and dynamic dataset switching.
    """
    filter_type = request.GET.get('period', 'month')
    from_date_raw = request.GET.get('from_date', '')
    to_date_raw = request.GET.get('to_date', '')
    active_tab = request.GET.get('tab', 'sales')

    start_date, end_date = get_date_range(filter_type, from_date_raw, to_date_raw)

    # Filter sales within resolved range
    sales_qs = Sale.objects.select_related('customer', 'user').prefetch_related('items__product__category').filter(
        date_sold__gte=start_date,
        date_sold__lte=end_date
    ).order_by('-date_sold', '-created_at')

    # Aggregated metrics for the filtered period
    sales_summary = sales_qs.aggregate(
        total_revenue=Sum('total_amount'),
        total_profit=Sum('total_profit'),
        total_cogs=Sum('total_cost'),
        total_sales_count=Count('id')
    )
    total_rev = sales_summary['total_revenue'] or Decimal('0.00')
    total_prof = sales_summary['total_profit'] or Decimal('0.00')
    total_cost = sales_summary['total_cogs'] or Decimal('0.00')

    # Filter expenses within resolved range
    expenses_qs = Expense.objects.select_related('category').filter(
        date__gte=start_date,
        date__lte=end_date
    ).order_by('-date', '-created_at')
    total_exp = expenses_qs.aggregate(tot=Sum('amount'))['tot'] or Decimal('0.00')
    net_operating_profit = total_prof - total_exp

    total_items_sold = SaleItem.objects.filter(sale__in=sales_qs).aggregate(tot=Sum('quantity'))['tot'] or 0

    # Current inventory valuation metrics
    active_batches = StockBatch.objects.filter(is_active=True, quantity_remaining__gt=0)
    inventory_val = sum((b.cost_price * b.quantity_remaining for b in active_batches), Decimal('0.00'))
    
    in_stock_products = Product.objects.select_related('category', 'stock').filter(stock__quantity_available__gt=0).order_by('category__name', 'name')
    expected_sales_val = sum((p.current_selling_price * p.available_quantity for p in in_stock_products), Decimal('0.00'))
    in_stock_count = in_stock_products.count()

    # Margin percentage
    margin_pct = (total_prof / total_rev * 100) if total_rev > 0 else Decimal('0.00')

    all_products = Product.objects.select_related('category', 'stock').order_by('category__name', 'name')

    return render(request, 'reports/index.html', {
        'filter_type': filter_type,
        'start_date': start_date,
        'end_date': end_date,
        'active_tab': active_tab,
        'total_rev': total_rev,
        'total_cost': total_cost,
        'total_prof': total_prof,
        'total_exp': total_exp,
        'net_operating_profit': net_operating_profit,
        'total_items_sold': total_items_sold,
        'margin_pct': margin_pct,
        'inventory_val': inventory_val,
        'expected_sales_val': expected_sales_val,
        'in_stock_count': in_stock_count,
        'sales': sales_qs[:50],  # Show recent 50 in table preview
        'expenses': expenses_qs[:50],
        'products': all_products[:50],
        'settings': CompanySetting.get_settings(),
    })



# ==============================================================================
# EXCEL EXPORT (openpyxl)
# ==============================================================================

def export_sales_excel(request):
    """Export filtered sales report to stylized Excel (.xlsx)."""
    filter_type = request.GET.get('period', 'month')
    start_date, end_date = get_date_range(filter_type, request.GET.get('from_date'), request.GET.get('to_date'))

    sales_qs = Sale.objects.select_related('customer', 'user').prefetch_related('items__product__category').filter(
        date_sold__gte=start_date,
        date_sold__lte=end_date
    ).order_by('date_sold')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Report"

    title_font = Font(name='Calibri', size=16, bold=True, color='1E293B')
    subtitle_font = Font(name='Calibri', size=10, italic=True, color='64748B')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    total_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    bold_font = Font(name='Calibri', size=11, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    settings_obj = CompanySetting.get_settings()

    # Document Header
    ws['A1'] = settings_obj.company_name.upper()
    ws['A1'].font = title_font
    ws['A2'] = f"Sales & Revenue Report ({start_date.strftime('%d/%m/%Y')} - {end_date.strftime('%d/%m/%Y')})"
    ws['A2'].font = subtitle_font

    # Column Headers
    headers = ['Date', 'Invoice #', 'Product(s)', 'Qty', 'Customer', 'Cost ($)', 'Total Sold ($)', 'Profit ($)', 'Payment']
    row_num = 4
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center' if col_num in (1, 2, 4, 9) else 'left')
        cell.border = thin_border

    total_rev = Decimal('0.00')
    total_cost = Decimal('0.00')
    total_profit = Decimal('0.00')

    for sale in sales_qs:
        row_num += 1
        items_desc = ", ".join([f"{it.product.name} (x{it.quantity})" + (f" [SN: {it.serial_number}]" if it.serial_number else "") for it in sale.items.all()])
        ws.cell(row=row_num, column=1, value=sale.date_sold.strftime('%Y-%m-%d')).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=2, value=sale.invoice_number).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=3, value=items_desc)
        ws.cell(row=row_num, column=4, value=sale.item_count).alignment = Alignment(horizontal='center')
        ws.cell(row=row_num, column=5, value=f"{sale.customer.name} ({sale.customer.phone})")
        
        cost_cell = ws.cell(row=row_num, column=6, value=float(sale.total_cost))
        cost_cell.number_format = '$#,##0.00'
        
        sell_cell = ws.cell(row=row_num, column=7, value=float(sale.total_amount))
        sell_cell.number_format = '$#,##0.00'
        
        prof_cell = ws.cell(row=row_num, column=8, value=float(sale.total_profit))
        prof_cell.number_format = '$#,##0.00'
        
        ws.cell(row=row_num, column=9, value=sale.get_payment_method_display()).alignment = Alignment(horizontal='center')

        for c in range(1, 10):
            ws.cell(row=row_num, column=c).border = thin_border

        total_cost += sale.total_cost
        total_rev += sale.total_amount
        total_profit += sale.total_profit

    # Summary Totals Row
    row_num += 1
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=5)
    total_label = ws.cell(row=row_num, column=1, value="TOTALS")
    total_label.font = bold_font
    total_label.alignment = Alignment(horizontal='right')

    tot_cost_cell = ws.cell(row=row_num, column=6, value=float(total_cost))
    tot_cost_cell.font = bold_font
    tot_cost_cell.number_format = '$#,##0.00'

    tot_rev_cell = ws.cell(row=row_num, column=7, value=float(total_rev))
    tot_rev_cell.font = bold_font
    tot_rev_cell.number_format = '$#,##0.00'

    tot_prof_cell = ws.cell(row=row_num, column=8, value=float(total_profit))
    tot_prof_cell.font = bold_font
    tot_prof_cell.number_format = '$#,##0.00'

    for c in range(1, 10):
        cell = ws.cell(row=row_num, column=c)
        cell.fill = total_fill
        cell.border = thin_border

    # Auto-fit Column Widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Sales_Report_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def export_inventory_excel(request):
    """Export complete current inventory valuation to Excel (.xlsx)."""
    products = Product.objects.select_related('category', 'stock').all().order_by('category__name', 'name')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventory Valuation"

    title_font = Font(name='Calibri', size=16, bold=True, color='1E293B')
    subtitle_font = Font(name='Calibri', size=10, italic=True, color='64748B')
    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='0D9488', end_color='0D9488', fill_type='solid')
    total_fill = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    bold_font = Font(name='Calibri', size=11, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    settings_obj = CompanySetting.get_settings()
    ws['A1'] = settings_obj.company_name.upper()
    ws['A1'].font = title_font
    ws['A2'] = f"Complete Physical Inventory Audit & Valuation — Generated {timezone.now().strftime('%d/%m/%Y %H:%M')}"
    ws['A2'].font = subtitle_font

    headers = ['Product Name', 'Category', 'Available Qty', 'Selling Price ($)', 'Expected Value ($)']
    row_num = 4
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=row_num, column=col_num, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center' if col_num == 3 else 'left')
        cell.border = thin_border

    tot_units = 0
    tot_valuation = Decimal('0.00')

    for p in products:
        row_num += 1
        qty = p.available_quantity
        sp = p.current_selling_price
        val = sp * qty
        ws.cell(row=row_num, column=1, value=p.name)
        ws.cell(row=row_num, column=2, value=p.category.name if p.category else '—')
        ws.cell(row=row_num, column=3, value=qty).alignment = Alignment(horizontal='center')
        
        c_sell = ws.cell(row=row_num, column=4, value=float(sp))
        c_sell.number_format = '$#,##0.00'

        c_val = ws.cell(row=row_num, column=5, value=float(val))
        c_val.number_format = '$#,##0.00'

        for c in range(1, 6):
            ws.cell(row=row_num, column=c).border = thin_border

        tot_units += qty
        tot_valuation += val

    # Totals
    row_num += 1
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=2)
    t_label = ws.cell(row=row_num, column=1, value="TOTAL INVENTORY VALUATION")
    t_label.font = bold_font
    t_label.alignment = Alignment(horizontal='right')

    ws.cell(row=row_num, column=3, value=tot_units).font = bold_font
    ws.cell(row=row_num, column=3).alignment = Alignment(horizontal='center')

    ws.cell(row=row_num, column=5, value=float(tot_valuation)).font = bold_font
    ws.cell(row=row_num, column=5).number_format = '$#,##0.00'

    for c in range(1, 6):
        cell = ws.cell(row=row_num, column=c)
        cell.fill = total_fill
        cell.border = thin_border

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 15)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Inventory_Valuation_{timezone.now().strftime('%Y%m%d')}.xlsx"
    response = HttpResponse(
        output.getvalue(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


# ==============================================================================
# PDF REPORTS (ReportLab Offline Engine)
# ==============================================================================

def export_sales_pdf(request):
    """Generate executive sales summary PDF."""
    filter_type = request.GET.get('period', 'month')
    start_date, end_date = get_date_range(filter_type, request.GET.get('from_date'), request.GET.get('to_date'))

    sales_qs = Sale.objects.select_related('customer', 'user').prefetch_related('items__product').filter(
        date_sold__gte=start_date,
        date_sold__lte=end_date
    ).order_by('date_sold')

    settings_obj = CompanySetting.get_settings()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )

    elements = []
    styles = getSampleStyleSheet()

    header_para = Paragraph(
        f"<b>{settings_obj.company_name.upper()}</b><br/>"
        f"<font size='9' color='#64748b'>{settings_obj.tagline} &bull; {settings_obj.phone}</font><br/>"
        f"<font size='14' color='#2563eb'><b>EXECUTIVE SALES & REVENUE REPORT</b></font><br/>"
        f"<font size='9'>Reporting Period: {start_date.strftime('%d %b %Y')} to {end_date.strftime('%d %b %Y')}</font>",
        styles['Normal']
    )
    elements.append(header_para)
    elements.append(Spacer(1, 15))

    # Table Header
    table_data = [
        ['Date', 'Invoice #', 'Item(s) Description', 'Units', 'Customer', 'Cost', 'Amount Sold', 'Profit', 'Payment']
    ]

    tot_rev = Decimal('0.00')
    tot_cost = Decimal('0.00')
    tot_profit = Decimal('0.00')

    for s in sales_qs:
        items_desc = ", ".join([f"{it.product.name} (x{it.quantity})" for it in s.items.all()])[:35]
        table_data.append([
            s.date_sold.strftime('%d/%m/%Y'),
            s.invoice_number,
            items_desc,
            str(s.item_count),
            s.customer.name[:18],
            f"${s.total_cost:,.2f}",
            f"${s.total_amount:,.2f}",
            f"${s.total_profit:,.2f}",
            s.payment_method
        ])
        tot_cost += s.total_cost
        tot_rev += s.total_amount
        tot_profit += s.total_profit

    # Total Row
    table_data.append([
        'TOTALS', '', '', f"{sales_qs.count()} Sales", '',
        f"${tot_cost:,.2f}", f"${tot_rev:,.2f}", f"${tot_profit:,.2f}", ''
    ])

    report_table = Table(table_data, colWidths=[65, 80, 160, 45, 110, 65, 75, 65, 65])
    report_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2563eb')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (5, 0), (7, -1), 'RIGHT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#f1f5f9')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))

    elements.append(report_table)
    elements.append(Spacer(1, 15))
    elements.append(Paragraph(f"Generated offline on {timezone.now().strftime('%d %b %Y %H:%M')} &bull; 21 Void Technologies", styles['Italic']))

    doc.build(elements)
    buffer.seek(0)

    filename = f"Sales_Report_{start_date.strftime('%Y%m%d')}_{end_date.strftime('%Y%m%d')}.pdf"
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{filename}"'
    return response


def export_inventory_pdf(request):
    """Generate executive inventory valuation PDF."""
    products = Product.objects.select_related('category', 'stock').all().order_by('category__name', 'name')
    settings_obj = CompanySetting.get_settings()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30
    )

    elements = []
    styles = getSampleStyleSheet()

    header_para = Paragraph(
        f"<b>{settings_obj.company_name.upper()}</b><br/>"
        f"<font size='9' color='#64748b'>{settings_obj.tagline} &bull; {settings_obj.phone}</font><br/>"
        f"<font size='14' color='#0d9488'><b>PHYSICAL INVENTORY AUDIT & VALUATION</b></font><br/>"
        f"<font size='9'>Stock as of {timezone.now().strftime('%d %B %Y %H:%M')}</font>",
        styles['Normal']
    )
    elements.append(header_para)
    elements.append(Spacer(1, 15))

    table_data = [
        ['Product Name', 'Category', 'Available Qty', 'Selling Price', 'Expected Value']
    ]

    tot_units = 0
    tot_val = Decimal('0.00')

    for p in products:
        qty = p.available_quantity
        sp = p.current_selling_price
        val = sp * qty
        table_data.append([
            p.name[:35],
            p.category.name[:25] if p.category else '—',
            str(qty),
            f"${sp:,.2f}",
            f"${val:,.2f}"
        ])
        tot_units += qty
        tot_val += val

    table_data.append([
        'TOTAL INVENTORY VALUATION', '', str(tot_units), '', f"${tot_val:,.2f}"
    ])

    report_table = Table(table_data, colWidths=[240, 160, 100, 110, 120])
    report_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0d9488')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (2, 0), (4, -1), 'RIGHT'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e2e8f0')),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#f1f5f9')),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))

    elements.append(report_table)
    elements.append(Spacer(1, 15))
    elements.append(Paragraph(f"Generated offline on {timezone.now().strftime('%d %b %Y %H:%M')} &bull; 21 Void Technologies", styles['Italic']))

    doc.build(elements)
    buffer.seek(0)

    filename = f"Inventory_Valuation_{timezone.now().strftime('%Y%m%d')}.pdf"
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{filename}"'
    return response
